"""Clean Ranpak lease ("User fee billing") invoice backups into the A+ upload .txt.

Automates Rhea's manual steps:
  1. Drop the logo/header rows above the column headings.
  2. Remove lines with a zero Period Amount (avoids zero invoices in A+).
  3. Sort by Ship Date; machines shipped before 2022 use the Old Serial#.
  4. Save as tab-delimited text (same format Excel's "Text (Tab delimited)" writes).

Several invoices can be combined into one upload. If the invoice PDF sits next
to the Excel file (same invoice number), the cleaned total is checked against it.
Serial numbers are looked up in A+ the same way the import does (see aplus_check.py).
It also writes Ranpak_Invoice_<invoice>_<timestamp>.txt, the single order file for the new
ZORHOF/ZORDOF map (see order_file.py).

Usage:
    python ranpak_clean.py 90210494.xlsx [more.xlsx ...] [-o upload.txt]
"""
import argparse
import datetime as dt
import re
import sys
from pathlib import Path

import openpyxl

OLD_SERIAL_CUTOFF = dt.datetime(2022, 1, 1)

HEADERS = [
    "Contract#", "Customer#", "Customer Name", "Install Location",
    "Install Location Name", "Install Location Addr", "Install Location City",
    "Install Location State", "Install Location Cntry", "Old Item", "New Item",
    "Item Description", "Old Serial#", "Serial#", "Ship Date", "Std User Fee",
    "Grace Period", "Deviation", "Total Disc", "Period Amount", "Demo Machine",
]
COL = {name: i for i, name in enumerate(HEADERS)}


class InvoiceError(Exception):
    pass


def fmt(value):
    """Format a cell the way Excel's tab-delimited export does for Rhea's files."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return f"{value.month}/{value.day}/{value.year}"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, int):
        return str(value)
    s = str(value)
    if s.strip():  # all-space fields are left padded, matching Excel's export
        s = s.strip()
        if re.fullmatch(r"-?\d+\.\d+", s):  # numbers stored as text, e.g. "100.00"
            f = float(s)
            s = str(int(f)) if f.is_integer() else str(f)
    if any(ch in s for ch in ',"\n'):
        s = '"' + s.replace('"', '""') + '"'
    return s


def to_number(value, where):
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        raise InvoiceError(f"{where}: Period Amount '{value}' is not a number")


def read_invoice(path):
    ws = openpyxl.load_workbook(path, data_only=True).active

    header_row = None
    info = {}
    for row in ws.iter_rows(min_row=1, max_row=30):
        values = [c.value for c in row]
        label = str(values[0]).strip() if values[0] is not None else ""
        if label in ("Invoice Type", "Invoice#", "Contract Description", "Master Contract#"):
            info[label] = str(values[1]).strip() if values[1] is not None else ""
        if [str(v).strip() if v is not None else "" for v in values[:len(HEADERS)]] == HEADERS:
            header_row = row[0].row
            break
    if header_row is None:
        raise InvoiceError(f"{path.name}: couldn't find the column headings - has Ranpak changed the layout?")
    if info.get("Invoice Type", "").lower() != "user fee billing":
        raise InvoiceError(f"{path.name}: Invoice Type is '{info.get('Invoice Type')}', expected 'User fee billing'")

    invoice_no = info.get("Invoice#") or path.stem
    lines, warnings = [], []
    raw_count = 0
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        row = list(row[:len(HEADERS)]) + [None] * (len(HEADERS) - len(row))
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue
        raw_count += 1
        where = f"{path.name} contract {fmt(row[COL['Contract#']])}"
        amount = to_number(row[COL["Period Amount"]], where)
        lines.append({"row": row, "amount": amount, "ship": row[COL["Ship Date"]], "where": where})

    # Same order Rhea gets in Excel: sort by Period Amount, drop zeros, then
    # sort by Ship Date (Excel's sort is stable, so ties keep the amount order).
    # Zero lines can have a blank Ship Date, so dates are only checked after the drop.
    lines.sort(key=lambda l: l["amount"])
    lines = [l for l in lines if l["amount"] != 0]
    for l in lines:
        if not isinstance(l["ship"], dt.datetime):
            raise InvoiceError(f"{l['where']}: Ship Date '{l['ship']}' is not a date")
        if any(isinstance(v, str) and ("\n" in v or "\t" in v) for v in l["row"]):
            warnings.append(f"{l['where']}: a field contains a line break or tab, which may break the A+ import")
    lines.sort(key=lambda l: l["ship"])

    swapped = 0
    for l in lines:
        row = l["row"]
        if l["ship"] < OLD_SERIAL_CUTOFF:
            old = row[COL["Old Serial#"]]
            if old is None or not str(old).strip():
                warnings.append(f"{l['where']}: shipped {fmt(l['ship'])} (before 2022) but has no Old Serial# - kept Serial# {fmt(row[COL['Serial#']])}")
            else:
                row[COL["Serial#"]] = old
                swapped += 1
        if not fmt(row[COL["Serial#"]]).strip():
            warnings.append(f"{l['where']}: no serial number")

    total = sum(l["amount"] for l in lines)
    return {
        "path": path, "invoice": invoice_no, "period": info.get("Contract Description", ""),
        "raw_count": raw_count, "lines": lines, "swapped": swapped,
        "total": total, "warnings": warnings,
    }


def pdf_total(xlsx_path, invoice_no):
    """Return the 'Final amount' from the matching invoice PDF, if one is next to the Excel file."""
    candidates = [p for p in xlsx_path.parent.iterdir()
                  if p.suffix.lower() == ".pdf" and invoice_no in p.stem]
    if not candidates:
        return None, None
    try:
        import pypdf
    except ImportError:
        return candidates[0], "pypdf not installed"
    text = "\n".join(page.extract_text() or "" for page in pypdf.PdfReader(candidates[0]).pages)
    m = re.search(r"([\d,]+\.\d{2})\s*Final amount", text)
    if not m:
        return candidates[0], "couldn't find 'Final amount' on the PDF"
    return candidates[0], float(m.group(1).replace(",", ""))


def parse_overrides(values):
    overrides = {}
    for v in values:
        sn, sep, target = v.partition("=")
        cust, dash, shipto = target.rpartition("-")
        if not sep or not dash or not sn.strip() or not cust.strip() or not shipto.strip():
            raise InvoiceError(f"--override '{v}' should look like SERIAL=CUSTOMER-SHIPTO, e.g. 10007655=501113-1")
        overrides[sn.strip()] = (cust.strip(), shipto.strip())
    return overrides


# Layout agreed with Trey (2026-10-07): the run details (who, where, when) are
# repeated on every line so each record has the same fields, and the invoice
# number and timestamp are also in the file name.
ORDER_FILE_HEADER = ("order_number|customer|ship_to|line_seq|item|description|serial|amount|invoice_number"
                     "|UserAccount|SourceComputer|Timestamp")


def order_file_path(folder, invoices, stamp):
    """Ranpak_Invoice_<InvoiceNumber>_<TimeStamp>.txt.

    A combined upload is named after its first invoice plus how many more it holds
    (e.g. 90210494+19), which keeps the path short; every line still carries its
    own invoice number.
    """
    name = invoices[0]["invoice"] + (f"+{len(invoices) - 1}" if len(invoices) > 1 else "")
    return folder / f"Ranpak_Invoice_{name}_{stamp}.txt"


def write_order_file(invoices, folder, overrides):
    """Write the single pipe-delimited order file for the ZORHOF/ZORDOF map; return warnings."""
    import getpass
    import platform
    from order_file import build_order_lines, used_rp_numbers
    lines = [{"invoice": inv["invoice"],
              "serial": fmt(l["row"][COL["Serial#"]]).strip(),
              "description": str(l["row"][COL["Item Description"]] or "").strip(),
              "amount": l["amount"]}
             for inv in invoices for l in inv["lines"]]
    text, orders, unmatched = build_order_lines(lines, used_rp_numbers(), overrides)
    stamp = f"{dt.datetime.now():%Y%m%d%H%M%S}"
    run = "|".join(v.replace("|", " ") for v in (getpass.getuser(), platform.node(), stamp))
    path = order_file_path(folder, invoices, stamp)
    with open(path, "w", encoding="cp1252", newline="") as f:
        f.write("\r\n".join([ORDER_FILE_HEADER] + [f"{t}|{run}" for t in text]) + "\r\n")
    print(f"Wrote {path} ({len(orders)} orders, {orders[0][0]}-{orders[-1][0]})" if orders else f"Wrote {path} (no orders)")
    return [f"Serial# {l['serial']} ({l['invoice']}): no customer found - left out of the order file" for l in unmatched]


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="+", type=Path, help="Ranpak Excel backup file(s)")
    ap.add_argument("-o", "--output", type=Path, help="output .txt (default: next to the first file)")
    ap.add_argument("--skip-aplus", action="store_true",
                    help="don't check serial numbers against A+ (also skips the order file)")
    ap.add_argument("--override", action="append", default=[], metavar="SERIAL=CUSTOMER-SHIPTO",
                    help="bill a serial to this customer/ship-to in the order file (e.g. 10007655=501113-1)")
    args = ap.parse_args(argv)

    invoices, problems = [], []
    try:
        overrides = parse_overrides(args.override)
    except InvoiceError as e:
        problems.append(str(e))
    for path in args.files:
        try:
            invoices.append(read_invoice(path))
        except InvoiceError as e:
            problems.append(str(e))
    if problems:
        print("STOPPED - nothing was written:")
        for p in problems:
            print("  -", p)
        return 1

    # Serial #s are unique across invoices, so a repeat means something is wrong.
    seen, dupes = {}, []
    for inv in invoices:
        for l in inv["lines"]:
            sn = fmt(l["row"][COL["Serial#"]]).strip()
            if sn and sn in seen:
                dupes.append(f"Serial# {sn} appears on both {seen[sn]} and {inv['invoice']}")
            seen.setdefault(sn, inv["invoice"])

    # Look each serial up the way the A+ import will, so problems show up before upload
    # instead of on the error report afterwards.
    aplus_warnings = []
    if not args.skip_aplus:
        try:
            from aplus_check import check_serials
            for sn, problem in check_serials(seen).items():
                if sn in overrides:
                    problem = f"overridden to {'-'.join(overrides[sn])} (A+ lookup: {problem})"
                aplus_warnings.append(f"Serial# {sn} ({seen[sn]}): {problem}")
        except Exception as e:
            aplus_warnings.append(f"couldn't check serial numbers against A+ ({e}) - the A+ error report is the only check")

    if args.output:
        out = args.output
    elif len(invoices) == 1:
        out = args.files[0].with_name(f"{invoices[0]['invoice']}.txt")
    else:
        out = args.files[0].with_name(f"ranpak_upload_{dt.date.today():%Y%m%d}.txt")

    print(f"{'Invoice':<10} {'Period':<12} {'Rows':>5} {'Billed':>6} {'Old SN':>6} {'Total':>12}  PDF check")
    grand, mismatch = 0.0, False
    for inv in invoices:
        pdf, pdf_amt = pdf_total(inv["path"], inv["invoice"])
        if pdf is None:
            check = "no PDF found"
        elif isinstance(pdf_amt, str):
            check = pdf_amt
        elif abs(pdf_amt - inv["total"]) < 0.005:
            check = "matches"
        else:
            check = f"MISMATCH - PDF says ${pdf_amt:,.2f}"
            mismatch = True
        print(f"{inv['invoice']:<10} {inv['period'][:12]:<12} {inv['raw_count']:>5} {len(inv['lines']):>6} "
              f"{inv['swapped']:>6} {inv['total']:>12,.2f}  {check}")
        grand += inv["total"]
    print(f"{'TOTAL':<10} {'':<12} {'':>5} {sum(len(i['lines']) for i in invoices):>6} {'':>6} {grand:>12,.2f}")

    # A total that doesn't match the invoice must never reach A+, so stop before writing.
    if mismatch:
        print("\nSTOPPED - nothing was written: an invoice total doesn't match its PDF (see above).")
        print("Check the Excel backup against the invoice before trying again.")
        return 1

    text_lines = ["\t".join(HEADERS)]
    for inv in invoices:
        text_lines += ["\t".join(fmt(v) for v in l["row"]) for l in inv["lines"]]
    with open(out, "w", encoding="cp1252", newline="") as f:
        f.write("\r\n".join(text_lines) + "\r\n")
    print(f"\nWrote {out}")

    # The order file for the new ZORHOF/ZORDOF map. It needs A+ for the customer lookup
    # and the RP number check, so it's skipped along with the A+ check.
    order_warnings = []
    if not args.skip_aplus:
        try:
            order_warnings = write_order_file(invoices, out.parent, overrides)
        except Exception as e:
            order_warnings.append(f"couldn't build the order file ({e})")

    warnings = [w for inv in invoices for w in inv["warnings"]] + dupes + aplus_warnings + order_warnings
    if warnings:
        print("\nCHECK BEFORE UPLOADING:")
        for w in warnings:
            print("  -", w)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
